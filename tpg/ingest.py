"""File -> clauses. No LLM here. Heading detection with a page-level fallback."""
import re
from html.parser import HTMLParser
from pathlib import Path

from tpg.models import Clause

HEADING_RE = re.compile(r"^(\d+(?:\.\d+)*)\.?[ \t ]+([A-Z].*)$")
NUMBER_ONLY_RE = re.compile(r"^(\d+\.\d+(?:\.\d+)*)\.?[ \t ]*$")
MAX_HEADING_LEN = 90


class _Text(HTMLParser):
    BLOCK = {"p", "div", "h1", "h2", "h3", "h4", "h5", "h6", "li", "tr", "br", "table", "section"}

    def __init__(self):
        super().__init__()
        self.parts: list[str] = []
        self._skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style"):
            self._skip += 1
        if tag in self.BLOCK:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in ("script", "style"):
            self._skip -= 1
        if tag in self.BLOCK:
            self.parts.append("\n")

    def handle_data(self, data):
        if not self._skip:
            self.parts.append(data)


def read_pages(path: str) -> list[str]:
    p = Path(path)
    suffix = p.suffix.lower()
    if suffix == ".pdf":
        import fitz
        with fitz.open(path) as doc:
            return [page.get_text() for page in doc]
    if suffix == ".docx":
        from docx import Document
        doc = Document(path)
        lines = [para.text for para in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                lines.append(" | ".join(cell.text for cell in row.cells))
        return ["\n".join(lines)]
    if suffix in (".html", ".htm"):
        parser = _Text()
        parser.feed(p.read_text(errors="replace"))
        return ["".join(parser.parts)]
    # Markdown, .txt, anything else: plain text. Strip leading '#' so headings match.
    text = p.read_text(errors="replace")
    return ["\n".join(re.sub(r"^#+\s*", "", line) for line in text.splitlines())]


def _heading(line: str) -> tuple[str, str] | None:
    if len(line) >= MAX_HEADING_LEN:
        return None
    m = HEADING_RE.match(line)
    if not m:
        return None
    number, title = m.group(1), m.group(2).strip()
    if title.lower().startswith("of "):  # "3 of 77" page footers
        return None
    return number, title


def _join_number_only_lines(lines: list[str]) -> list[str]:
    """ISO PDFs print '7.4.2\\t' on one line and the title on the next."""
    out: list[str] = []
    i = 0
    while i < len(lines):
        line = lines[i]
        if NUMBER_ONLY_RE.match(line.strip()) and i + 1 < len(lines) and lines[i + 1].strip():
            out.append(line.strip().rstrip(".") + " " + lines[i + 1].strip())
            i += 2
        else:
            out.append(line)
            i += 1
    return out


def segment(pages: list[str]) -> list[Clause]:
    clauses: list[Clause] = []
    current: Clause | None = None
    for page in pages:
        for raw in _join_number_only_lines(page.splitlines()):
            line = raw.replace(" ", " ").strip()
            head = _heading(line)
            if head:
                current = Clause(id=head[0], title=head[1], text="")
                clauses.append(current)
            elif current is not None and line:
                current.text = (current.text + "\n" + line).strip()
    if clauses:
        return clauses
    return [Clause(id=f"p{n}", title="", text=page.strip()) for n, page in enumerate(pages, 1)]


def ingest(path: str) -> list[Clause]:
    return segment(read_pages(path))
