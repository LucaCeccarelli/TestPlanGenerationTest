from pathlib import Path

import pytest

from tpg.ingest import CONTROL, follows, ingest, segment, strip_repeated, table_rows

FIX = Path(__file__).parent / "fixtures"


@pytest.mark.parametrize("name", ["sample.md", "sample.html", "sample.docx", "sample.pdf"])
def test_sample_segments_into_known_clauses(name):
    clauses = ingest(str(FIX / name))
    ids = [c.id for c in clauses]
    expected = ["1", "5.1", "5.2", "5.3", "5.4"]
    assert expected == [i for i in ids if i in set(expected)]
    c51 = next(c for c in clauses if c.id == "5.1")
    assert c51.title == "Response time"
    assert "within 500 ms" in c51.text
    c52 = next(c for c in clauses if c.id == "5.2")
    assert "shall not process" in c52.text
    assert "Name: nonce | Presence: mandatory" in next(c for c in clauses if c.id == "5.4").text


def test_html_ignores_script():
    text = " ".join(c.text + c.title for c in ingest(str(FIX / "sample.html")))
    assert "var x" not in text


def test_iso_style_number_on_own_line_is_joined():
    pages = ["7.4.2\t\nBiometric template\nThe template shall be present.\n7.5\t Country codes\nCodes shall follow ISO 3166."]
    clauses = segment(pages)
    assert [(c.id, c.title) for c in clauses] == [("7.4.2", "Biometric template"), ("7.5", "Country codes")]


def test_nbsp_after_number_is_accepted():
    clauses = segment(["5.1.\xa0\xa0New Parameters\nThe wallet shall accept them."])
    assert clauses[0].id == "5.1" and clauses[0].title == "New Parameters"


def test_page_footer_is_not_a_heading():
    clauses = segment(["3 of 77\nSome body text that shall be kept."])
    assert all(c.id != "3" for c in clauses)


def test_no_headings_falls_back_to_pages():
    clauses = segment(["first page text", "second page text"])
    assert [c.id for c in clauses] == ["p1", "p2"]
    assert clauses[1].text == "second page text"


def test_long_line_is_not_a_heading():
    line = "5 " + "x" * 100
    clauses = segment([line])
    assert clauses[0].id == "p1"


def test_openid_pdf_has_expected_clauses():
    clauses = ingest(str(FIX / "OpenID4VP1-0.pdf"))
    ids = {c.id for c in clauses}
    assert {"5", "5.1", "5.2"} <= ids
    assert {"9", "10", "11", "12", "13"} <= ids
    assert max(len(c.text) for c in clauses) < 60000


def test_duplicate_ids_are_unique_after_segmentation():
    pages = ["1 New Parameters\n1 New Parameters\nText A.\n2 Other\nBody other.\n1 New Parameters\nText B."]
    clauses = segment(pages)
    ids = [c.id for c in clauses]
    assert len(ids) == len(set(ids))
    assert ids == ["1", "2", "1#2"]
    assert next(c for c in clauses if c.id == "1").text == "Text A."


def test_fullest_duplicate_keeps_bare_id():
    # a bare-id transition backward (5.7 -> 5.8 -> 5.7) is rejected by follows(); a ToC-then-body
    # duplicate reaches the same id again only via the "1" restart, as real ToCs do.
    pages = ["5.7 Title\nshort\n7 Other\nx\n1 Intro\n3 More\n5 Overview\n"
             "5.1 Sub\nBody one.\n5.3 Sub2\nBody three.\n5.5 Sub3\nBody five.\n"
             "5.7 Title\nThe unit shall do the long real thing here.\n5.9 Next\ny"]
    clauses = segment(pages)
    by_id = {c.id: c for c in clauses}
    assert by_id["5.7"].text.startswith("The unit shall") and by_id["5.7#2"].text == "short"


def test_openid_real_clauses_keep_bare_ids():
    by_id = {c.id: c for c in ingest(str(FIX / "OpenID4VP1-0.pdf"))}
    assert len(by_id["5.7"].text) > 200 and "5.7#2" in by_id


def test_openid_pdf_ids_are_unique():
    clauses = ingest(str(FIX / "OpenID4VP1-0.pdf"))
    ids = [c.id for c in clauses]
    assert len(ids) == len(set(ids))


def test_unnumbered_markdown_falls_back_to_one_clause_per_heading(tmp_path):
    md = tmp_path / "u.md"
    md.write_text("# Intro\nText one.\n## Rules\nThe unit shall work.\n## Notes\nMore.\n")
    clauses = ingest(str(md))
    assert [c.id for c in clauses] == ["p1", "p2", "p3"]
    assert "shall work" in clauses[1].text


def test_openid_pdf_text_has_no_control_characters():
    text = "\n".join(c.title + "\n" + c.text for c in ingest(str(FIX / "OpenID4VP1-0.pdf")))
    assert not any(ord(ch) < 32 and ch not in "\n\t\r" for ch in text)
    assert "Verifier" in text


def test_docx_tables_stay_in_document_order(tmp_path):
    from docx import Document
    doc = Document()
    doc.add_paragraph("5.1 First")
    t = doc.add_table(rows=2, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Item", "Value"
    t.rows[1].cells[0].text, t.rows[1].cells[1].text = "The unit shall log", "events"
    doc.add_paragraph("5.2 Second")
    doc.add_paragraph("Body two.")
    p = tmp_path / "t.docx"
    doc.save(str(p))
    clauses = ingest(str(p))
    assert "Item: The unit shall log | Value: events" in next(c for c in clauses if c.id == "5.1").text
    assert "shall log" not in next(c for c in clauses if c.id == "5.2").text


def test_rfc_pdf_falls_back_to_pages():
    clauses = ingest(str(FIX / "RFC8949.pdf"))
    assert clauses[0].id == "p1" and len(clauses) > 50


@pytest.mark.skipif(not (FIX / "iso_18013_5.pdf").exists(), reason="ISO fixture not distributed")
def test_iso_pdf_has_expected_clauses():
    ids = {c.id for c in ingest(str(FIX / "iso_18013_5.pdf"))}
    assert {"1", "6.1", "8.1", "8.1.1"} <= ids


def test_strip_repeated_removes_running_headers_and_page_numbers():
    pages = [f"Spec v1.0\nBody {i}\nmore text {i}\nPage {i} of 4" for i in range(1, 5)]
    out = strip_repeated(pages)
    assert out[0] == "Body 1\nmore text 1"
    assert all("Spec v1.0" not in p and "Page" not in p for p in out)


def test_strip_repeated_keeps_short_documents_and_mid_page_lines():
    pages = ["Spec v1.0\nBody\nSpec v1.0 again mid", "Spec v1.0\nOther"]
    assert strip_repeated(pages) == pages
    four = ["Spec\n" + "\n".join(f"line {j}" for j in range(10)) + "\nTable 3 caption\nend"] * 4
    assert all("Table 3 caption" in p for p in strip_repeated(four))


def test_control_characters_do_not_hide_repeated_footers():
    pages = [f"Body {i}\nline a\nline b\nline c\nline d\n(c) Publisher 2020 - All rights reserved{chr(8) if i % 2 else ''}\n{i}" for i in range(1, 7)]
    out = strip_repeated([p.translate(CONTROL) for p in pages])
    assert all("All rights reserved" not in p for p in out)


def test_iso_pdf_has_no_footer_inside_clauses():
    if not (FIX / "iso_18013_5.pdf").exists():
        pytest.skip("ISO fixture not distributed")
    clauses = ingest(str(FIX / "iso_18013_5.pdf"))
    hits = sum(c.text.count("All rights reserved") for c in clauses)
    assert hits <= 2, hits
    assert len(clauses) >= 250


def test_table_rows_format_and_merged_first_cell():
    rows = [["Name", "Presence", "Type"], ["nonce", "M", "text"], ["", "O", "int"], ["", "", ""]]
    assert table_rows(rows) == "Name: nonce | Presence: M | Type: text\nName: nonce | Presence: O | Type: int"


def test_single_row_table_keeps_its_text(tmp_path):
    md = tmp_path / "t.md"
    md.write_text("## 5.4 Note\n\n| The unit shall log every event. | see 5.1 |\n")
    assert "The unit shall log every event. | see 5.1" in ingest(str(md))[0].text


def test_follows_sequence_rules():
    assert follows(None, "1") and follows("1", "2") and follows("1", "1.1") and follows("1.1", "1.2")
    assert follows("7.4.9", "7.5") and follows("7.4.9", "8") and follows("1.1", "1.3")  # one missed heading tolerated
    assert follows("8", "8.1") and not follows("8", "8.1.2") and not follows("1.1", "1.5")
    assert follows("B.2", "B.3") and follows("B", "B.1") and follows("A.3.4", "B") and not follows("B.2", "5.2")
    assert follows("9.3.2", "A") and follows("11", "A") and not follows("E.2.2", "10") and not follows("17", "8949")
    assert follows("5", "6.1") and follows("7.2", "7.3.1")
    assert follows("9", "1") and follows("C.3", "1", prev_empty=True) and not follows("8.6", "1")
    assert not follows("8.6", "1", prev_empty=False)


def test_letter_prefixed_and_annex_headings():
    clauses = segment(["1 Scope\nBody.\nAnnex A Use cases\nIntro.\nA.1 General\nThe unit shall log.\nA.2 Cases\nMore."])
    assert [c.id for c in clauses] == ["1", "A", "A.1", "A.2"]
    assert clauses[1].title == "Use cases"


def test_table_cell_number_is_not_a_heading():
    clauses = segment(["B.1 Extensions\nAuthority Key Identifier\n5.2 Further extensions shall not be present\nCRL Number\nB.2 Next\nBody."])
    assert [c.id for c in clauses] == ["B.1", "B.2"]
    assert "5.2 Further extensions" in clauses[0].text


def test_joined_heading_accepts_lowercase_title_but_inline_lowercase_is_rejected():
    clauses = segment(["7.3.2\t\nnameSpace\nThe nameSpace shall be text.\n7.3.3 mDL data\nBody."])
    assert [(c.id, c.title) for c in clauses] == [("7.3.2", "nameSpace")]
    assert "7.3.3 mDL data" in clauses[0].text


def test_soft_hyphens_are_removed():
    clauses = segment(["5.1 Codes\nThe first part of the code shall be the same as issu\u00ading_country."])
    assert "issuing_country" in clauses[0].text


def test_markdown_pipe_table_becomes_row_records(tmp_path):
    md = tmp_path / "t.md"
    md.write_text("## 5.4 Parameters\n\n| Name | Presence | Type |\n|---|---|---|\n| nonce | mandatory | text string |\n| locale | optional | language tag |\n")
    c = ingest(str(md))[0]
    assert "Name: nonce | Presence: mandatory | Type: text string" in c.text
    assert "Name: locale | Presence: optional | Type: language tag" in c.text
    assert "|---" not in c.text


def test_html_table_becomes_row_records(tmp_path):
    h = tmp_path / "t.html"
    h.write_text("<h2>5.4 Parameters</h2><table><tr><th>Name</th><th>Presence</th></tr><tr><td>nonce</td><td>mandatory</td></tr></table>")
    c = ingest(str(h))[0]
    assert "Name: nonce | Presence: mandatory" in c.text


def test_docx_table_becomes_row_records_with_headers(tmp_path):
    from docx import Document
    doc = Document()
    doc.add_paragraph("5.4 Parameters")
    t = doc.add_table(rows=2, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Name", "Presence"
    t.rows[1].cells[0].text, t.rows[1].cells[1].text = "nonce", "mandatory"
    p = tmp_path / "t.docx"
    doc.save(str(p))
    c = next(c for c in ingest(str(p)) if c.id == "5.4")
    assert "Name: nonce | Presence: mandatory" in c.text


def test_iso_pdf_finds_lowercase_titled_and_annex_clauses():
    if not (FIX / "iso_18013_5.pdf").exists():
        pytest.skip("ISO fixture not distributed")
    clauses = {c.id: c for c in ingest(str(FIX / "iso_18013_5.pdf"))}
    assert {"7.3.3", "A.1", "8.2.1.1.2.1"} <= set(clauses)
    assert max(len(c.text) for c in clauses.values()) < 60000
    assert "Reference:" in "".join(c.text for c in clauses.values())   # a table row record survived


def test_pdf_pages_have_no_running_headers():
    text = "\n".join(c.text for c in ingest(str(FIX / "OpenID4VP1-0.pdf")))
    assert " of 96" not in text
