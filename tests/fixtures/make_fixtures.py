"""Regenerate sample.docx, sample.pdf and sample.html from sample.md. Run once; outputs are committed."""
import re
import textwrap
from pathlib import Path

import pymupdf
from docx import Document

from tpg.ingest import table_rows   # the row-record format the ingest side reads back

HERE = Path(__file__).parent
lines = (HERE / "sample.md").read_text().splitlines()

# Group the Markdown into ("heading"|"para", text) and ("table", rows) blocks.
blocks: list[tuple[str, object]] = []
for line in lines:
    if line.startswith("|"):
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if all(set(c) <= set("-: ") for c in cells):   # the |---|---| separator
            continue
        if blocks and blocks[-1][0] == "table":
            blocks[-1][1].append(cells)
        else:
            blocks.append(("table", [cells]))
    elif line.startswith("#"):
        blocks.append(("heading", line))
    elif line.strip():
        blocks.append(("para", line))

# DOCX: headings become paragraphs "5.1 Response time", body stays as is, tables become real tables.
doc = Document()
for kind, value in blocks:
    if kind == "heading":
        doc.add_paragraph(value.lstrip("#").strip())
    elif kind == "para":
        doc.add_paragraph(value)
    else:
        t = doc.add_table(rows=len(value), cols=len(value[0]))
        for row, cells in zip(t.rows, value):
            for cell, text in zip(row.cells, cells):
                cell.text = text
doc.save(HERE / "sample.docx")

# PDF: mimic ISO layout: the clause number on its own line, title on the next.
# Drawn table borders are not reliably detected, so tables go in as row-record text lines.
pdf = pymupdf.open()
page = pdf.new_page()
y = 40
for kind, value in blocks:
    if kind == "heading":
        m = re.match(r"#+\s+(\d+(?:\.\d+)*)\s+(.*)", value)
        if m and "." in m.group(1):          # ISO layout: dotted number alone, title on the next line
            page.insert_text((40, y), m.group(1)); y += 14
            page.insert_text((40, y), m.group(2)); y += 14
            continue
        page.insert_text((40, y), value.lstrip("#").strip()); y += 14
    else:
        text = value if kind == "para" else table_rows(value)
        for line in text.splitlines():
            for chunk in textwrap.wrap(line, 100):  # keep every line inside the page or get_text() drops it
                page.insert_text((40, y), chunk, fontsize=9); y += 14
pdf.save(HERE / "sample.pdf")

# HTML: headings as <h2>, paragraphs as <p>, tables as <table>, plus a script that must be ignored.
html = ["<html><head><title>Sample</title><script>var x = 1;</script></head><body>"]
for kind, value in blocks:
    if kind == "heading":
        html.append(f"<h2>{value.lstrip('#').strip()}</h2>")
    elif kind == "para":
        html.append(f"<p>{value}</p>")
    else:
        rows = ["<tr>" + "".join(f"<th>{c}</th>" for c in value[0]) + "</tr>"]
        rows += ["<tr>" + "".join(f"<td>{c}</td>" for c in cells) + "</tr>" for cells in value[1:]]
        html.append("<table>" + "".join(rows) + "</table>")
html.append("</body></html>")
(HERE / "sample.html").write_text("\n".join(html))
print("fixtures written")
